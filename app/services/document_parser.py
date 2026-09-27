"""
FoundRisk Edge — document parsing.

Dispatches by file type to the right library (PyMuPDF / pandas / openpyxl),
extracts text/table content, and hands segments to claim_extractor. Scanned PDF
pages fall back to local Tesseract OCR when the Tesseract binary is installed.
"""
import fitz  # PyMuPDF
import openpyxl
import pandas as pd
import io
import os
import shutil

try:
    import pytesseract
    from PIL import Image
except ImportError:  # OCR remains optional for text-PDF-only installs.
    pytesseract = None
    Image = None

from app.services import claim_extractor as extractor


class ParseResult:
    def __init__(self, status, status_detail=None, page_count=None, claims=None, ocr_pages=0):
        self.status = status                # "ready" | "failed" | "empty"
        self.status_detail = status_detail  # human-readable message, or None
        self.page_count = page_count
        self.claims = claims or []
        self.ocr_pages = ocr_pages


# Minimum characters of extracted text per page before we trust it's a real
# text-based PDF and not a scanned/image-only one.
_MIN_TEXT_CHARS_PER_PAGE = 20


def _configure_tesseract():
    """Use an explicit Tesseract path when configured, otherwise PATH lookup."""
    if pytesseract is None:
        return False
    configured = os.environ.get("TESSERACT_CMD", "").strip()
    if configured:
        pytesseract.pytesseract.tesseract_cmd = configured
    elif not shutil.which("tesseract"):
        # Common Windows install location; no path is assumed if it is absent.
        candidate = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if os.path.isfile(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
        else:
            return False
    return True


def _ocr_pdf_page(page):
    """Render one PDF page and OCR it locally; return (text, unavailable)."""
    if not _configure_tesseract():
        return "", True
    try:
        scale = float(os.environ.get("FOUNDRISK_OCR_SCALE", "2"))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        language = os.environ.get("FOUNDRISK_OCR_LANG", "eng")
        return pytesseract.image_to_string(image, lang=language, config="--psm 6", timeout=45), False
    except Exception as exc:
        # A missing executable is an installation problem; other OCR errors are
        # reported as an attempted OCR with no readable text, not as success.
        not_found = getattr(getattr(pytesseract, "pytesseract", None), "TesseractNotFoundError", ())
        return "", bool(not_found and isinstance(exc, not_found))


def parse_pdf(filepath: str) -> ParseResult:
    try:
        doc = fitz.open(filepath)
    except Exception as exc:  # malformed/corrupt PDF
        return ParseResult(status="failed", status_detail=f"Could not open PDF: {exc}")

    if doc.page_count == 0:
        return ParseResult(status="empty", status_detail="PDF contains no pages.", page_count=0)

    page_count = doc.page_count
    page_texts = []
    ocr_pages = 0
    ocr_unavailable = False
    for page_index in range(page_count):
        page = doc.load_page(page_index)
        text = page.get_text("text") or ""
        if len(text.strip()) < _MIN_TEXT_CHARS_PER_PAGE:
            ocr_text, unavailable = _ocr_pdf_page(page)
            ocr_unavailable = ocr_unavailable or unavailable
            if len(ocr_text.strip()) > len(text.strip()):
                text = ocr_text
                if text.strip():
                    ocr_pages += 1
        page_texts.append(text)
    doc.close()

    total_text_chars = sum(len(text.strip()) for text in page_texts)
    if total_text_chars < _MIN_TEXT_CHARS_PER_PAGE * page_count * 0.5:
        if ocr_unavailable:
            detail = ("Scanned/image PDF detected, but local OCR is unavailable. Install Tesseract OCR and "
                      "the English language data, then retry. See README.md → Scanned PDF OCR.")
        else:
            detail = "Scanned/image PDF detected. OCR was attempted locally, but no readable text was found."
        return ParseResult(status="empty", status_detail=detail, page_count=page_count, ocr_pages=ocr_pages)

    all_claims = []
    for page_index, text in enumerate(page_texts, start=1):
        if len(text.strip()) < _MIN_TEXT_CHARS_PER_PAGE:
            continue
        all_claims.extend(extractor.extract_from_pdf_page(text, page_index))

    detail = f"Local OCR extracted text from {ocr_pages} scanned page(s)." if ocr_pages else None
    return ParseResult(status="ready", status_detail=detail, page_count=page_count, claims=all_claims, ocr_pages=ocr_pages)


def parse_csv(filepath: str) -> ParseResult:
    try:
        df = pd.read_csv(filepath, dtype=str, keep_default_na=False)
    except Exception as exc:
        return ParseResult(status="failed", status_detail=f"Could not read CSV: {exc}")

    if df.empty:
        return ParseResult(status="empty", status_detail="CSV contains no data rows.", page_count=0)

    all_claims = []
    for idx, row in df.iterrows():
        row_number = idx + 2  # +1 for 0-index, +1 for header row
        row_dict = {col: row[col] for col in df.columns}
        all_claims.extend(extractor.extract_from_csv_row(row_dict, row_number))

    return ParseResult(status="ready", page_count=len(df), claims=all_claims)


def parse_xlsx(filepath: str) -> ParseResult:
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    except Exception as exc:
        return ParseResult(status="failed", status_detail=f"Could not read XLSX: {exc}")

    sheet_names = wb.sheetnames
    if not sheet_names:
        return ParseResult(status="empty", status_detail="Workbook contains no sheets.", page_count=0)

    all_claims = []
    any_data = False

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=False)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            continue
        headers = [str(cell.value) if cell.value is not None else f"col_{i}" for i, cell in enumerate(header_row)]

        sheet_rows = []
        for row_cells in rows_iter:
            row_has_data = any(cell.value not in (None, "") for cell in row_cells)
            if not row_has_data:
                continue
            any_data = True
            # Each column maps to {"value": ..., "cell": "B14"} so the actual
            # worksheet cell coordinate is preserved for provenance — never
            # reconstructed from header name + row number.
            row_dict = {}
            for header, cell in zip(headers, row_cells):
                row_dict[header] = {
                    "value": cell.value if cell.value is not None else "",
                    "cell": cell.coordinate,
                }
            sheet_rows.append(row_dict)

        all_claims.extend(extractor.extract_from_xlsx_sheet(sheet_name, sheet_rows))

    wb.close()

    if not any_data:
        return ParseResult(status="empty", status_detail="Workbook contains no data rows.", page_count=len(sheet_names))

    return ParseResult(status="ready", page_count=len(sheet_names), claims=all_claims)


def parse_document(filepath: str, file_type: str) -> ParseResult:
    if file_type == "pdf":
        return parse_pdf(filepath)
    if file_type == "csv":
        return parse_csv(filepath)
    if file_type == "xlsx":
        return parse_xlsx(filepath)
    return ParseResult(status="failed", status_detail=f"Unsupported file type: {file_type}")
